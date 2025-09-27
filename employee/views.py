from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.decorators import login_required
from .models import Employee, Country, JobTitle, UserProfile
from .forms import EmployeeForm, RegisterForm
from .decorators import role_required
from django.core.paginator import Paginator

import openpyxl
from django.contrib import messages



# Create your views here.

def user_login(request):
    if request.method == 'POST':
        username = request.POST['username']
        password = request.POST['password']
        user = authenticate(username=username, password=password)
        
        if user is not None:
            login(request, user)
            return redirect('dashboard')
        else:
            messages.error(request,'Invalid Username or Password')
            return redirect('login')
    return render(request, 'login.html')


def user_logout(request):
    logout(request)
    return render(request, "login.html")



@login_required
def dashboard(request):
    user_role = request.user.userprofile.role
    if user_role == 'SUPER_ADMIN':
        employees = Employee.objects.all()
    elif user_role in ['HR_MANAGER', 'HR_EXECUTIVE']:
        employees = Employee.objects.all()
    else:
        employees = [get_object_or_404(Employee, user=request.user)]
    if isinstance(employees, list):
        paginated_employees = employees 
        page_obj = None
    else:
        paginator = Paginator(employees, 12) 
        page_number = request.GET.get('page')
        paginated_employees = paginator.get_page(page_number)
        page_obj = paginated_employees 

    return render(request, 'dashboard.html', {
        'employees': paginated_employees,
        'user_role': user_role,
        'page_obj': page_obj
    })



@login_required
@role_required(['HR_EXECUTIVE','SUPER_ADMIN'])

def add_employee(request):
    if request.method == 'POST':
        form = EmployeeForm(request.POST, request.FILES)
        if form.is_valid():
            employee = form.save(commit=False)
            employee.user = request.user
            employee.save()
            messages.success(request, "Employee added successfully!")
            return redirect('dashboard')
    else:
        form = EmployeeForm()
    return render(request, 'add_employee.html', {'form': form, 'countries': Country.objects.all(), 'jobs': JobTitle.objects.all()})

@login_required
@role_required(['HR_MANAGER','SUPER_ADMIN'])

def edit_employee(request, pk):
    employee = get_object_or_404(Employee, pk=pk)
    if request.method == 'POST':
        form = EmployeeForm(request.POST,request.FILES, instance=employee)
        if form.is_valid():
            form.save()
            messages.success(request, "Employee updated successfully!")
            return redirect('dashboard')
    else:
        form = EmployeeForm(instance=employee)
    return render(request, 'edit_employee.html', {'employee': employee, 'countries': Country.objects.all(), 'jobs': JobTitle.objects.all()})



@login_required
@role_required(['HR_MANAGER','SUPER_ADMIN'])

def delete_employee(request, pk):
    employee = get_object_or_404(Employee, pk=pk)
    employee.delete()
    return redirect('dashboard')

def register(request):
    if request.method == 'POST':
        form = RegisterForm(request.POST)
        if form.is_valid():
            user = form.save(commit=False)
            user.email = form.cleaned_data['email']
            user.save()

            role = form.cleaned_data['role']
            UserProfile.objects.create(user=user, role=role)

            Employee.objects.create(
                user=user,
                name=form.cleaned_data['name'],
                mobile=form.cleaned_data['mobile'],
                phone=form.cleaned_data['phone'],
                country=form.cleaned_data['country'],
                city=form.cleaned_data['city'],
                job=form.cleaned_data['job'],
                profile_photo=form.cleaned_data['profile_photo']
            )
            return redirect('login')
    else:
        form = RegisterForm()
    return render(request, 'register.html', {'form': form})

@login_required
def employee_table(request):
    employees = Employee.objects.all()
    return render(request, 'table_employee.html', {'employees': employees})

@login_required
def import_excel(request):
    if request.method == "POST":
        excel_file = request.FILES['excel_file']
        wb = openpyxl.load_workbook(excel_file)
        sheet = wb.active

        expected_headers = ['name', 'mobile', 'phone', 'country', 'city', 'job']
        actual_headers = [str(cell.value).strip().lower() for cell in next(sheet.iter_rows(min_row=1, max_row=1))]

        if actual_headers != expected_headers:
            messages.error(request, "Excel headers do not match the required format. Please use the sample file.")
            return redirect('dashboard')
        
        for row in sheet.iter_rows(min_row=2, values_only=True):
            name, mobile, phone, country, city, job = row
            
            if name is None:
                name = 'N/A'
            if mobile is None:
                mobile = 'N/A'
            if phone is None:
                phone = 'N/A'
            if country is None:
                country = 'N/A'
            if city is None:
                city = 'N/A'
            if job is None:
                job = 'N/A'
            
            phone = str(phone)[:20] if phone else None 


            country_obj, created = Country.objects.get_or_create(name=country)
            job_obj, created = JobTitle.objects.get_or_create(title=job)

            Employee.objects.create(
                user=request.user,
                name=name,
                mobile=mobile,
                phone=phone,
                country=country_obj,
                city=city,
                job=job_obj
            )
        messages.success(request, "Employees imported successfully.")
    return redirect('dashboard')

@login_required
def employee_table(request):
    user_role = request.user.userprofile.role
    employees = Employee.objects.all()
    paginator = Paginator(employees, 12)
    page_number = request.GET.get('page')
    page_obj = paginator.get_page(page_number)

    return render(request, 'table_employee.html', {
        'employees': page_obj,
        'user_role': user_role,
        'page_obj': page_obj
    })


@login_required
def search_emp(request):
    user_role = request.user.userprofile.role
    query = request.GET.get('q')
    if query:
        employees = Employee.objects.filter(name__icontains=query)
        if employees.count() == 1:
            contexts ={
                'employees': employees,
                'user_role': user_role,
                'is_single_employee': True,
                
            }
            return render(request, 'dashboard.html',contexts)
        else:
            contexts ={
                'employees': employees,
                'user_role': user_role,
                'is_single_employee': False,
                
            }
            return render(request,'dashboard.html',contexts)
    else:
        if user_role in ['SUPER_ADMIN','HR_MANAGER','HR_EXECUTIVE']:
            employees = Employee.objects.all()
        else:
            employees = [get_object_or_404(Employee, user=request.user)]

    return render(request, 'dashboard.html', {
        'employees': employees,
        'user_role': user_role,
        'is_single_employee': False,
        
    })
    